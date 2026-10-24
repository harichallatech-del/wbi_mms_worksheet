"""Dry-run-first submitter for queues embedded in the role workbook screens."""
from __future__ import annotations

import argparse
import json
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests

from wbi_common import WBIError, configure_logging, fail_with_log
from wbi_role_layout import QUEUE_RANGE, ROLE_ENV, ROLE_SHEETS, role_from_book_sheets

SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]
ALLOWED = re.compile(
    r"^/(?:manufacturing-(?:bw|hb)|stock-positions|machine-maintenance|admin/(?:status|products|operators|machines|auth-users))(?:/[A-Za-z0-9_-]+)?$"
    r"|^/raw-materials/(?:by-name|[A-Za-z0-9_-]+)$|^/stock-levels$"
    r"|^/contact-inquiries/[A-Za-z0-9_-]+/(?:status|whatsapp)$"
)

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--platform", choices=("excel", "google"), required=True)
    parser.add_argument("--file", help="Local role workbook XLSX (Excel mode).")
    parser.add_argument("--role", choices=("manager", "stock-monitor", "operators"), help="Required for Google mode; Excel detects the role from the tab list.")
    parser.add_argument("--commit", action="store_true", help="Submit queue entries; omission is a dry run.")
    parser.add_argument("--sync-role-users", action="store_true", help="Refresh active account names in lookup dropdowns from WBI-MMS.")
    return parser.parse_args()

def google_service():
    credential = os.getenv("GOOGLE_APPLICATION_CREDENTIALS", "").strip()
    if not credential or not Path(credential).is_file():
        raise WBIError("Google mode needs GOOGLE_APPLICATION_CREDENTIALS pointing to a service-account key outside this project.")
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build
    creds = Credentials.from_service_account_file(credential, scopes=SCOPES)
    return build("sheets", "v4", credentials=creds, cache_discovery=False)

def _excel_context(args: argparse.Namespace):
    if not args.file:
        raise WBIError("Excel mode requires --file with one of the role workbooks.")
    from openpyxl import load_workbook
    path = Path(args.file)
    if not path.is_file():
        raise WBIError(f"Workbook not found: {path}")
    wb = load_workbook(path)
    role = role_from_book_sheets(wb.sheetnames)
    return role, (wb, path)

def _google_id(role: str) -> str:
    env = ROLE_ENV[role]
    value = os.getenv(env, "").strip()
    if not value:
        raise WBIError(f"Set {env} to the imported Google Sheet ID for the {role} workbook.")
    return value

def read_queue(args: argparse.Namespace):
    if args.platform == "excel":
        role, ctx = _excel_context(args)
        wb, _ = ctx
        rows = []
        for sheet in ROLE_SHEETS[role]:
            ws = wb[sheet]
            for row in range(2, 2002):
                values = [ws.cell(row, col).value for col in range(45, 54)]
                if values[0] and str(values[6] or "").upper() == "PENDING":
                    rows.append({"sheet": sheet, "row": row, "values": values})
        return role, rows, ctx
    role = args.role
    if not role:
        raise WBIError("Google mode requires --role so the correct separate role workbook is selected.")
    service, sid = google_service(), _google_id(role)
    meta = service.spreadsheets().get(spreadsheetId=sid, fields="sheets.properties.title").execute()
    names = [s["properties"]["title"] for s in meta.get("sheets", [])]
    if set(names) != set(ROLE_SHEETS[role]):
        raise WBIError(f"Google Sheet tabs {names} do not match the {role} role screen list.")
    rows = []
    for sheet in ROLE_SHEETS[role]:
        result = service.spreadsheets().values().get(spreadsheetId=sid, range=f"'{sheet}'!{QUEUE_RANGE}", valueRenderOption="UNFORMATTED_VALUE").execute()
        for row, values in enumerate(result.get("values", []), start=2):
            values += [""] * (9 - len(values))
            if values[0] and str(values[6]).upper() == "PENDING":
                rows.append({"sheet": sheet, "row": row, "values": values[:9]})
    return role, rows, (service, sid)

def api_request(session: requests.Session, base: str, token: str, method: str, path: str, payload: dict[str, Any] | None = None):
    if not ALLOWED.fullmatch(path):
        raise WBIError("Queue path is outside the WBI form endpoint allowlist.")
    response = session.request(method, base.rstrip("/") + path, headers={"Authorization": f"Bearer {token}", "Accept": "application/json", "Content-Type": "application/json"}, json=payload, timeout=(10, 45))
    return response

def resolve_payload(session: requests.Session, base: str, token: str, path: str, payload: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    if path in ("/raw-materials/by-name", "/stock-levels") and "raw_material_name" in payload:
        response = api_request(session, base, token, "GET", "/stock-levels")
        if not response.ok:
            raise WBIError(f"Material lookup failed ({response.status_code}).")
        body = response.json(); rows = body if isinstance(body, list) else body.get("data", body.get("items", [])) if isinstance(body, dict) else []
        material = next((x for x in rows if str(x.get("raw_material_name", x.get("name", ""))).casefold() == str(payload["raw_material_name"]).casefold()), None)
        if not material or not material.get("raw_material_id", material.get("id")):
            raise WBIError("Could not resolve raw material name to a WBI ID.")
        payload = dict(payload); material_id = material.get("raw_material_id", material.get("id")); payload.pop("raw_material_name", None)
        if path == "/stock-levels":
            return f"/raw-materials/{material_id}", payload
        return f"/raw-materials/{material_id}", payload
    if path == "/machine-maintenance" and "machine_name" in payload:
        response = api_request(session, base, token, "GET", "/admin/machines")
        if not response.ok:
            raise WBIError(f"Machine lookup failed ({response.status_code}).")
        body = response.json(); rows = body if isinstance(body, list) else body.get("data", body.get("items", [])) if isinstance(body, dict) else []
        machine = next((x for x in rows if str(x.get("name", x.get("machine_name", ""))).casefold() == str(payload["machine_name"]).casefold()), None)
        if not machine or not machine.get("machine_id", machine.get("id")):
            raise WBIError("Could not resolve machine name to a WBI ID.")
        payload = dict(payload); payload["machine_id"] = machine.get("machine_id", machine.get("id")); payload.pop("machine_name", None)
    return path, payload

def update_row(args: argparse.Namespace, ctx: Any, item: dict[str, Any], status: str, message: str, safe_payload: dict[str, Any] | None = None) -> None:
    sheet, row = item["sheet"], item["row"]
    updated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    if args.platform == "excel":
        wb, _ = ctx; ws = wb[sheet]
        ws.cell(row, 51, status); ws.cell(row, 52, message[:240]); ws.cell(row, 53, updated)
        if safe_payload is not None: ws.cell(row, 50, json.dumps(safe_payload, ensure_ascii=False))
        return
    service, sid = ctx
    service.spreadsheets().values().update(spreadsheetId=sid, range=f"'{sheet}'!AY{row}:BA{row}", valueInputOption="RAW", body={"values": [[status, message[:240], updated]]}).execute()
    if safe_payload is not None:
        service.spreadsheets().values().update(spreadsheetId=sid, range=f"'{sheet}'!AX{row}", valueInputOption="RAW", body={"values": [[json.dumps(safe_payload, ensure_ascii=False)]]}).execute()

def normalize_role_users(body: Any) -> tuple[list[str], list[str], list[list[Any]]]:
    rows = body if isinstance(body, list) else body.get("data", body.get("items", body.get("users", []))) if isinstance(body, dict) else []
    stock, operators, manager_rows = [], [], []
    for user in rows:
        if not isinstance(user, dict): continue
        active = user.get("is_active", True)
        is_active = not (active is False or active == 0 or str(active).strip().lower() in ("false", "no", "inactive", "0"))
        name, role = str(user.get("name", "")).strip(), str(user.get("role", "")).strip().lower()
        if not name: continue
        manager_rows.append([name, str(user.get("phone", "")), role, "Yes" if is_active else "No"])
        if not is_active: continue
        target = stock if role == "stock-monitor" else operators if role == "operator" else None
        if target is not None and name.casefold() not in {x.casefold() for x in target}: target.append(name)
    return stock[:54], operators[:54], manager_rows[:68]

def _update_excel_roster(args: argparse.Namespace, ctx: Any, stock: list[str], operators: list[str], manager_rows: list[list[Any]]) -> None:
    wb, path = ctx
    for sheet in ROLE_SHEETS[role_from_book_sheets(wb.sheetnames)]:
        ws = wb[sheet]
        for offset in range(54):
            ws.cell(offset + 2, 38).value = stock[offset] if offset < len(stock) else None
            ws.cell(offset + 2, 39).value = operators[offset] if offset < len(operators) else None
    if "Admin Settings" in wb.sheetnames:
        ws = wb["Admin Settings"]
        for offset in range(68):
            for col in range(1, 5): ws.cell(offset + 33, col).value = manager_rows[offset][col - 1] if offset < len(manager_rows) else None
    wb.save(path)

def _update_google_roster(service: Any, sid: str, role: str, stock: list[str], operators: list[str], manager_rows: list[list[Any]]) -> None:
    meta = service.spreadsheets().get(spreadsheetId=sid, fields="sheets.properties.title").execute()
    names = [s["properties"]["title"] for s in meta.get("sheets", [])]
    if set(names) != set(ROLE_SHEETS[role]): raise WBIError(f"Google Sheet tabs do not match {role} role screens.")
    updates = []
    for sheet in ROLE_SHEETS[role]:
        stock_col = [[stock[i] if i < len(stock) else ""] for i in range(54)]
        operator_col = [[operators[i] if i < len(operators) else ""] for i in range(54)]
        updates.extend([
            {"range": f"'{sheet}'!AL2:AL55", "majorDimension": "ROWS", "values": stock_col},
            {"range": f"'{sheet}'!AM2:AM55", "majorDimension": "ROWS", "values": operator_col},
        ])
    if role == "manager":
        rows = manager_rows + [[""] * 4 for _ in range(68 - len(manager_rows))]
        updates.append({"range": "'Admin Settings'!A33:D100", "majorDimension": "ROWS", "values": rows})
    service.spreadsheets().values().batchUpdate(spreadsheetId=sid, body={"valueInputOption": "RAW", "data": updates}).execute()

def sync_role_users(args: argparse.Namespace) -> None:
    base, token = os.getenv("WBI_API_BASE_URL", "").strip(), os.getenv("WBI_API_TOKEN", "").strip()
    if not base or not token: raise WBIError("Account sync needs WBI_API_BASE_URL and an authorized Manager/Dev WBI_API_TOKEN.")
    response = api_request(requests.Session(), base, token, "GET", "/admin/auth-users")
    if not response.ok: raise WBIError(f"WBI account roster request failed ({response.status_code}); Manager/Dev access may be required.")
    stock, operators, manager_rows = normalize_role_users(response.json())
    if args.platform == "excel":
        role, ctx = _excel_context(args)
        _update_excel_roster(args, ctx, stock, operators, manager_rows)
    else:
        service = google_service()
        roles = [args.role] if args.role else list(ROLE_SHEETS)
        for role in roles: _update_google_roster(service, _google_id(role), role, stock, operators, manager_rows)
    print(f"Synced {len(stock)} Stock Monitor names and {len(operators)} Operator names. No account names were added to passwords/passkey data.")

def run() -> None:
    args = parse_args(); logger = configure_logging()
    if args.sync_role_users:
        sync_role_users(args); return
    role, records, ctx = read_queue(args)
    if not records: print("No PENDING actions found."); return
    base, token = os.getenv("WBI_API_BASE_URL", "").strip(), os.getenv("WBI_API_TOKEN", "").strip()
    session = requests.Session()
    for item in records:
        values = item["values"] + [""] * max(0, 9 - len(item["values"]))
        queue_id, _, resource, method, path, raw_payload, *_ = values
        try:
            payload = json.loads(raw_payload or "{}")
            if not isinstance(payload, dict): raise WBIError("Queue payload must be a JSON object.")
            method, path = str(method).upper(), str(path)
            if method not in ("POST", "PUT", "DELETE") or not ALLOWED.fullmatch(path): raise WBIError("Queue method or path is outside the WBI endpoint allowlist.")
            if path.startswith(("/admin/status", "/admin/products")): payload = {k: payload[k] for k in ("name", "description") if k in payload}
            elif path.startswith("/admin/operators"): payload = {k: payload[k] for k in ("name", "type") if k in payload}
            elif path.startswith("/admin/machines"): payload = {k: payload[k] for k in ("name", "type", "location") if k in payload}
            elif path == "/admin/auth-users":
                payload = {k: payload[k] for k in ("name", "phone", "role", "is_active") if k in payload}
                if not payload.get("name") or not payload.get("phone") or payload.get("role") not in ("stock-monitor", "operator", "manager"):
                    raise WBIError("Login-user creation needs a name, mobile number, and Manager, Stock Monitor, or Operator role.")
                payload["is_active"] = bool(payload.get("is_active", True))
            if not args.commit:
                print(f"DRY RUN {queue_id}: {method} {path} ({resource}); payload keys={','.join(sorted(payload)) or 'none'}")
                continue
            if not base or not token: raise WBIError("Set WBI_API_BASE_URL and an authorized WBI_API_TOKEN to commit actions.")
            resolved_path, resolved_payload = resolve_payload(session, base, token, path, payload)
            response = api_request(session, base, token, method, resolved_path, None if method == "DELETE" else resolved_payload)
            if not response.ok:
                message = f"WBI rejected request ({response.status_code}); check role permissions and required fields."
                update_row(args, ctx, item, "ERROR", message)
                logger.error("Queue item %s failed (%s)", queue_id, response.status_code, extra={"stage": "queue_submit"})
                print(f"ERROR {queue_id}: {message}")
            else:
                safe = {"name": payload.get("name", ""), "role": payload.get("role", "")} if path == "/admin/auth-users" else None
                update_row(args, ctx, item, "COMPLETED", f"HTTP {response.status_code}", safe_payload=safe)
                logger.info("Queue item %s completed (%s)", queue_id, response.status_code, extra={"stage": "queue_submit"})
                print(f"COMPLETED {queue_id}: HTTP {response.status_code}")
        except Exception as exc:
            message = str(exc).replace("\n", " ")[:220]
            if args.commit:
                try: update_row(args, ctx, item, "ERROR", message)
                except Exception: pass
            logger.error("Queue item %s failed: %s", queue_id, message, extra={"stage": "queue_submit"})
            print(f"ERROR {queue_id}: {message}")
    if args.platform == "excel" and args.commit:
        ctx[0].save(ctx[1])

def main() -> None:
    try: run()
    except Exception as exc: raise SystemExit(fail_with_log("queue_submit", exc)) from None

if __name__ == "__main__": main()
